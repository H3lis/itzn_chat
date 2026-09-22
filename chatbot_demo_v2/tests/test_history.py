"""대화 이력 및 만족도 통계 서비스 단위 테스트 (test_history.py)."""
import pytest
from pathlib import Path
from chatbot_demo_v2.app.history_service import HistoryService
from chatbot_demo_v2.config.settings import load_settings


@pytest.fixture
def history_service(tmp_path):
    settings = load_settings()
    return HistoryService(settings, db_path=tmp_path / "chat_history.db")


def test_record_and_search_history(history_service):
    # 1. 턴 기록 (개인정보 포함 질의)
    res = history_service.record_turn(
        session_id="sess_001",
        run_id="run_001",
        raw_question="교무실 홍길동 교사 PC IP 192.168.1.50 장애 발생. 010-1234-5678 연락요망",
        final_answer="네트워크 설정 가이드를 확인하세요.",
        route="faq",
        route_reason="FAQ 1위 매칭",
        latency_s=0.12,
        confidence="high",
        source_meta={"type": "faq", "id": "faq_015"},
    )
    assert res["has_pii"] is True
    assert "010-****-5678" in res["masked_question"]
    assert "192.168.*.*" in res["masked_question"]
    assert "홍**" in res["masked_question"]

    # 2. 검색
    search_res = history_service.search_history(page=1, page_size=10)
    assert search_res["total"] == 1
    item = search_res["items"][0]
    assert item["run_id"] == "run_001"
    assert item["route"] == "faq"
    assert item["feedback"] == "NONE"
    assert "010-****-5678" in item["masked_question"]

    # 3. 개인정보보호법 준수 검증: 원본 질문(raw_question) 및 원본 개인정보가 DB에 남지 않음
    detail = history_service.get_turn_detail("run_001")
    assert detail is not None
    assert "raw_question" not in detail or not detail.get("raw_question")
    assert "홍길동" not in str(detail.values())
    assert "010-1234-5678" not in str(detail.values())
    assert "192.168.1.50" not in str(detail.values())


def test_feedback_and_analytics(history_service):
    # 3개 레코드 생성
    history_service.record_turn(
        session_id="s1",
        run_id="r1",
        raw_question="스쿨넷 장애",
        final_answer="조치 방법 1",
        route="faq",
        latency_s=0.1,
    )
    history_service.record_turn(
        session_id="s2",
        run_id="r2",
        raw_question="AP 전원 불량",
        final_answer="조치 방법 2",
        route="rag3x",
        latency_s=0.5,
    )
    history_service.record_turn(
        session_id="s3",
        run_id="r3",
        raw_question="인터넷 연결 문의",
        final_answer="조치 방법 3",
        route="web_search",
        latency_s=1.2,
    )

    # 피드백 업데이트
    assert history_service.update_feedback("r1", "POSITIVE") is True
    assert history_service.update_feedback("r2", "NEGATIVE", reason="설명이 부족함") is True

    # 단건 상세 조회
    detail = history_service.get_turn_detail("r2")
    assert detail is not None
    assert detail["feedback"] == "NEGATIVE"
    assert detail["feedback_reason"] == "설명이 부족함"

    # 통계 요약 조회
    analytics = history_service.get_analytics_summary()
    assert analytics["total_queries"] == 3
    assert analytics["positive_feedback"] == 1
    assert analytics["negative_feedback"] == 1
    assert analytics["satisfaction_rate"] == 50.0  # 1 / 2 * 100
    assert len(analytics["recent_negatives"]) == 1
    assert analytics["recent_negatives"][0]["run_id"] == "r2"


def test_export_history_excel(history_service):
    import openpyxl

    # 1. 테스트 데이터 적재
    history_service.record_turn(
        session_id="sess_excel_1",
        run_id="run_ex_1",
        raw_question="홍길동 선생님 010-9876-5432 AP 고장 문의",
        final_answer="AP 전원을 재인가하세요.",
        route="faq",
        latency_s=0.25,
    )
    history_service.record_turn(
        session_id="sess_excel_2",
        run_id="run_ex_2",
        raw_question="인터넷 속도가 너무 느려요",
        final_answer="스위치 케이블을 점검하세요.",
        route="rag3x",
        latency_s=0.85,
    )
    history_service.update_feedback("run_ex_1", "POSITIVE")
    history_service.update_feedback("run_ex_2", "NEGATIVE", reason="개선안 불명확")

    # 2. 전체 내보내기 검증
    buf = history_service.export_history_excel()
    assert buf is not None
    wb = openpyxl.load_workbook(buf)
    assert "대화상담이력" in wb.sheetnames
    ws = wb["대화상담이력"]

    # 헤더 행 검증
    headers = [cell.value for cell in ws[1]]
    assert headers == [
        "번호",
        "상담일시",
        "세션 ID",
        "처리 경로",
        "사용자 질문 (비식별화)",
        "챗봇 응답",
        "소요시간(초)",
        "만족도",
        "피드백 사유",
        "비식별화 항목",
    ]
    # 데이터 행 수: 헤더 1행 + 데이터 2행 = 3행
    assert ws.max_row == 3

    # 첫 번째 행 값 검증 (최신순이므로 run_ex_2가 먼저)
    row2 = [cell.value for cell in ws[2]]
    assert row2[0] == 1
    assert row2[2] == "sess_excel_2"
    assert row2[3] == "RAG 심층검색"
    assert "인터넷 속도" in row2[4]
    assert "👎 불만족" in row2[7]
    assert row2[8] == "개선안 불명확"

    # 두 번째 행 값 검증 (run_ex_1)
    row3 = [cell.value for cell in ws[3]]
    assert row3[0] == 2
    assert row3[2] == "sess_excel_1"
    assert row3[3] == "FAQ 매칭"
    assert "홍**" in row3[4]
    assert "010-****-5432" in row3[4]
    assert "👍 만족" in row3[7]
    assert "인명" in row3[9] or "전화번호" in row3[9]

    # 3. 필터 조건(feedback='POSITIVE') 지정 내보내기 검증
    buf_pos = history_service.export_history_excel(feedback="POSITIVE")
    wb_pos = openpyxl.load_workbook(buf_pos)
    ws_pos = wb_pos["대화상담이력"]
    assert ws_pos.max_row == 2  # 헤더 1 + 데이터 1건
    assert "👍 만족" in ws_pos.cell(row=2, column=8).value


def test_export_history_excel_api(tmp_path):
    from fastapi.testclient import TestClient
    from chatbot_demo_v2.app.main import create_app

    app = create_app()
    client = TestClient(app)

    # API 호출
    resp = client.get("/api/admin/history/export/excel?feedback=all")
    assert resp.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in resp.headers["content-type"]
    assert "attachment; filename=" in resp.headers["content-disposition"]
    assert len(resp.content) > 0

