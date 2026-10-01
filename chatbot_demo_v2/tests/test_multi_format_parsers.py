# -*- coding: utf-8 -*-
"""다양한 문서 포맷(XLSX, DOCX, TXT 등) 파싱 및 색인 단위 테스트."""
import tempfile
from pathlib import Path
import pandas as pd
import pytest

from chatbot_demo_v2.ragcore.rag3.config import load_config
from chatbot_demo_v2.ragcore.rag3.converters import (
    parse_excel_file,
    parse_docx_file,
    parse_text_file,
)
from chatbot_demo_v2.ragcore.rag3.parse import parse_document


def test_excel_markdown_table_parsing():
    """엑셀 파일이 시트별 마크다운 표로 정확히 변환되는지 테스트."""
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        excel_path = tmp_path / "test_manual.xlsx"

        # 2개 시트 생성
        with pd.ExcelWriter(excel_path, engine="openpyxl") as writer:
            df1 = pd.DataFrame({
                "장비명": ["Switch-A", "Router-B"],
                "IP주소": ["192.168.1.1", "10.0.0.1"],
                "장애조치": ["재부팅", "케이블 점검"]
            })
            df1.to_excel(writer, sheet_name="네트워크장비", index=False)

            df2 = pd.DataFrame({
                "담당자": ["홍길동", "이순신"],
                "부서": ["인프라팀", "보안팀"],
                "내선번호": ["1001", "1002"]
            })
            df2.to_excel(writer, sheet_name="비상연락망", index=False)

        config = load_config(str(Path(__file__).resolve().parents[1] / "ragcore" / "rag3" / "config.yaml"))
        doc_info = parse_excel_file(excel_path, "test_manual.xlsx", config)

        assert doc_info.document_name == "test_manual.xlsx"
        assert doc_info.page_count == 2
        assert doc_info.parser_used == "excel_markdown_table"

        # 첫 번째 페이지 확인
        p1 = doc_info.pages[0]
        assert "네트워크장비" in p1.text
        assert "| Switch-A | 192.168.1.1 | 재부팅 |" in p1.text
        assert p1.has_table is True
        assert p1.page_type == "table"

        # 두 번째 페이지 확인
        p2 = doc_info.pages[1]
        assert "비상연락망" in p2.text
        assert "| 홍길동 | 인프라팀 | 1001 |" in p2.text


def test_docx_parsing():
    """DOCX 파일이 텍스트 및 표를 정확히 파싱하는지 테스트."""
    import docx

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        docx_path = tmp_path / "test_guide.docx"

        doc = docx.Document()
        doc.add_heading("네트워크 장애 가이드", level=1)
        doc.add_paragraph("인터넷 연결이 끊겼을 때의 대응 절차입니다.")

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "단계"
        table.cell(0, 1).text = "조치사항"
        table.cell(1, 0).text = "1단계"
        table.cell(1, 1).text = "공유기 전원 리셋"
        doc.save(str(docx_path))

        config = load_config(str(Path(__file__).resolve().parents[1] / "ragcore" / "rag3" / "config.yaml"))
        doc_info = parse_docx_file(docx_path, "test_guide.docx", config)

        assert doc_info.document_name == "test_guide.docx"
        assert doc_info.page_count >= 1
        assert "네트워크 장애 가이드" in doc_info.pages[0].text
        assert "공유기 전원 리셋" in doc_info.pages[0].text


def test_parse_document_dispatcher():
    """parse_document 함수가 확장자에 따라 올바른 파서로 디스패치하는지 테스트."""
    with tempfile.TemporaryDirectory() as tmpdir:
        tmp_path = Path(tmpdir)
        txt_path = tmp_path / "sample.txt"
        txt_path.write_text("스쿨넷 회선 장애 긴급 점검 요령\n\n모뎀 LED 상태를 확인하세요.", encoding="utf-8")

        config = load_config(str(Path(__file__).resolve().parents[1] / "ragcore" / "rag3" / "config.yaml"))
        doc_info = parse_document(txt_path, "sample.txt", config)

        assert doc_info.document_name == "sample.txt"
        assert doc_info.parser_used == "text_direct"
        assert "스쿨넷 회선 장애 긴급 점검 요령" in doc_info.pages[0].text
