# -*- coding: utf-8 -*-
"""엑셀 모범 질답 데이터를 data/faq.json 으로 정규화한다.

raw_data/faq/ 폴더 내 엑셀 파일을 읽어 data/faq.json을 생성합니다.

실행:
    python chatbot_demo_v2/scripts/import_faq.py
    python chatbot_demo_v2/scripts/import_faq.py --excel path/to/faq.xlsx
"""

from __future__ import annotations

import argparse
import json
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

from openpyxl import load_workbook

_PKG_ROOT = Path(__file__).resolve().parents[1]
if str(_PKG_ROOT.parent) not in sys.path:
    sys.path.insert(0, str(_PKG_ROOT.parent))

from chatbot_demo_v2.scenario.matcher import normalize_text  # noqa: E402

DEFAULT_FAQ_DIR = _PKG_ROOT / "raw_data" / "faq"
DEFAULT_OUT = _PKG_ROOT / "data" / "faq.json"
EXPECTED_SHEETS = ["스쿨넷", "학내망", "무선망", "유무선통합관제"]
_SPLIT_CHARS = ["\n", ";", ",", "/"]


def _find_default_excel() -> Path:
    """raw_data/faq/ 아래 첫 번째 .xlsx 파일을 찾거나 기본 경로 반환."""
    if DEFAULT_FAQ_DIR.is_dir():
        excels = [f for f in DEFAULT_FAQ_DIR.glob("*.xlsx") if not f.name.startswith("~$")]
        if excels:
            return excels[0]
    return DEFAULT_FAQ_DIR / "장애_상담_데이터.xlsx"


def _norm_header(h) -> str:
    """헤더 정규화: 공백 제거 후 비교용."""
    if h is None:
        return ""
    return str(h).replace(" ", "").strip()


def _clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def _split_source_files(raw) -> list[str]:
    if raw is None:
        return []
    text = str(raw).strip()
    if not text:
        return []
    parts = [text]
    for ch in _SPLIT_CHARS:
        nxt: list[str] = []
        for p in parts:
            nxt.extend(p.split(ch))
        parts = nxt
    seen: list[str] = []
    for p in parts:
        p = p.strip()
        if p and p not in seen:
            seen.append(p)
    return seen


def _column_map(header_row) -> dict[str, int]:
    """정규화 헤더명 → 0-based 컬럼 인덱스."""
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        key = _norm_header(cell)
        if key and key not in mapping:
            mapping[key] = idx
    return mapping


def import_excel(excel_path: Path, out_path: Path) -> dict:
    if not excel_path.is_file():
        raise SystemExit(f"[import_faq] 엑셀 파일을 찾을 수 없습니다: {excel_path}")

    wb = load_workbook(excel_path, read_only=True, data_only=True)
    entries: list[dict] = []
    per_sheet: dict[str, int] = {}
    skipped_empty = 0
    seen_norm: dict[str, str] = {}
    dup_warnings: list[str] = []

    for sheet_name in EXPECTED_SHEETS:
        if sheet_name not in wb.sheetnames:
            print(f"[import_faq] 주의: 시트 '{sheet_name}'를 찾을 수 없어 건너뜁니다.")
            per_sheet[sheet_name] = 0
            continue
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration:
            per_sheet[sheet_name] = 0
            continue
        cmap = _column_map(header)

        c_no = cmap.get("No")
        c_qtype = cmap.get("질문유형")
        c_ftype = cmap.get("장애유형")
        c_q = cmap.get("질문")
        c_a = cmap.get("답변")
        c_src = cmap.get("질문답변생성근거파일명")

        if c_q is None or c_a is None:
            raise SystemExit(
                f"[import_faq] {sheet_name}: 질문/답변 컬럼을 찾지 못함. 헤더={header}"
            )

        count = 0
        for r_idx, row in enumerate(rows, start=2):
            question = _clean(row[c_q]) if c_q < len(row) else None
            answer = _clean(row[c_a]) if c_a < len(row) else None
            if not question or not answer:
                skipped_empty += 1
                continue

            no_val = None
            if c_no is not None and c_no < len(row):
                raw_no = row[c_no]
                try:
                    no_val = int(raw_no) if raw_no is not None else None
                except (ValueError, TypeError):
                    no_val = None

            qtype = _clean(row[c_qtype]) if c_qtype is not None and c_qtype < len(row) else None
            ftype = _clean(row[c_ftype]) if c_ftype is not None and c_ftype < len(row) else None
            src = row[c_src] if c_src is not None and c_src < len(row) else None

            norm_q = normalize_text(question)
            entry_id = f"{sheet_name}:{r_idx}"
            if norm_q in seen_norm:
                dup_warnings.append(
                    f"중복 정규화 질문: {entry_id} == {seen_norm[norm_q]} ({norm_q[:30]})"
                )
            else:
                seen_norm[norm_q] = entry_id

            entries.append(
                {
                    "id": entry_id,
                    "sheet": sheet_name,
                    "row": r_idx,
                    "no": no_val,
                    "question_type": qtype,
                    "fault_type": ftype,
                    "question": question,
                    "question_normalized": norm_q,
                    "answer": answer,
                    "source_files": _split_source_files(src),
                }
            )
            count += 1
        per_sheet[sheet_name] = count

    wb.close()

    kst = timezone(timedelta(hours=9))
    payload = {
        "version": 1,
        "generated_at": datetime.now(kst).isoformat(),
        "source_file": excel_path.name,
        "entry_count": len(entries),
        "per_sheet": per_sheet,
        "entries": entries,
    }

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8") as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)

    print(f"[import_faq] 총 {len(entries)}건 → {out_path}")
    for name in EXPECTED_SHEETS:
        print(f"   - {name}: {per_sheet.get(name, 0)}건")
    print(f"[import_faq] 빈 질문/답변 skip: {skipped_empty}건")
    if dup_warnings:
        print(f"[import_faq] 경고: 중복 정규화 질문 {len(dup_warnings)}건")
    return payload


def main() -> int:
    parser = argparse.ArgumentParser(description="엑셀 모범 질답 → faq.json")
    parser.add_argument("--excel", type=Path, default=None, help="FAQ 원본 엑셀 경로")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT, help="출력 JSON 경로")
    args = parser.parse_args()

    excel_path = args.excel or _find_default_excel()
    import_excel(excel_path, args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
