# -*- coding: utf-8 -*-
"""수정된 RAG 엔진으로 바탕화면의 전체질문_개선RAG_비교.xlsx 엑셀 파일의 모든 시트를 채운다."""
from __future__ import annotations

import os
import re
import sys
import time
import traceback
from pathlib import Path
import openpyxl

# 'chatbot_demo_v2'의 상위 폴더를 sys.path에 추가하여 패키지로 임포트 가능하도록 설정
PKG_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PKG_ROOT.parent))

from chatbot_demo_v2.config.settings import load_settings
from chatbot_demo_v2.app.dependencies import build_context
from chatbot_demo_v2.graph.nodes import _render_persona_prompts


def compose_rag_answer(ctx, question: str, base_answer: str, evidence_text: str) -> str:
    """RAG 초안과 근거를 바탕으로 composer_rag 프롬프트를 렌더링하고 합성 답변을 생성합니다."""
    prompt = ctx.prompts.render(
        "composer_rag",
        question=question,
        answer=base_answer,
        evidence_text=evidence_text,
        history_summary="",
    )

    persona_text, _ = _render_persona_prompts(ctx, "rag3x")
    if persona_text:
        prompt = f"{persona_text}\n\n{prompt}"

    for attempt in range(1, 4):
        try:
            return (ctx.llm.chat(prompt) or "").strip()
        except Exception as e:
            print(f"    [LLM 재시도 {attempt}/3] 에러: {e}")
            if attempt == 3:
                raise e
            time.sleep(5 * attempt)
    return ""


def process_question_with_retry(ctx, question: str):
    """RAG 질의 및 합성 답변을 재시도 로직과 함께 실행합니다."""
    for attempt in range(1, 4):
        try:
            res = ctx.rag_adapter.ask(question)
            base_answer = res.get("final_answer") or ""
            evidence_text = res.get("answer_context") or ""
            composed_answer = compose_rag_answer(ctx, question, base_answer, evidence_text)
            return res, composed_answer
        except Exception as e:
            print(f"    [RAG 질의 재시도 {attempt}/3] 에러: {e}")
            if attempt == 3:
                raise e
            time.sleep(5 * attempt)


def main():
    excel_path = Path("c:/Users/hgcha/Desktop/전체질문_개선RAG_비교.xlsx")
    if not excel_path.exists():
        print(f"[오류] 엑셀 파일을 찾을 수 없습니다: {excel_path}")
        return 1

    print(f"엑셀 파일 로딩 중: {excel_path}")
    wb = openpyxl.load_workbook(excel_path)
    
    print("RAG 환경 로딩 중...")
    settings = load_settings()
    
    # 평가용이므로 캐시 비활성화하여 정밀한 실시간 생성을 보장
    object.__setattr__(settings, "rag_cache_ttl_s", 0)
    
    ctx = build_context(settings)
    print("RAG 엔진 예열(Warmup) 시작...")
    ctx.rag_adapter.warm_up(deep=False)
    print("예열 완료!")

    # 정규식 패턴 정의 (대소문자 무시)
    summary_pat = re.compile(r"<summary>([\s\S]*?)</summary>", re.IGNORECASE)
    details_pat = re.compile(r"<details>([\s\S]*?)</details>", re.IGNORECASE)

    total_success = 0
    total_skipped = 0

    print(f"\n처리할 시트 목록: {wb.sheetnames}")

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        max_row = sheet.max_row
        print(f"\n{'='*50}")
        print(f"시트 시작: [{sheet_name}] (최대 행: {max_row})")
        print(f"{'='*50}")

        sheet_success = 0
        sheet_skipped = 0

        for r in range(2, max_row + 1):
            q_val = sheet.cell(row=r, column=4).value  # D열: 질문
            if not q_val or not str(q_val).strip():
                continue

            rag_sum_val = sheet.cell(row=r, column=7).value  # G열: RAG 답변 - 요약
            rag_det_val = sheet.cell(row=r, column=8).value  # H열: RAG 답변 - 상세
            rag_src_val = sheet.cell(row=r, column=9).value  # I열: RAG 답변 생성 근거 파일명

            # 요약과 상세가 이미 모두 채워져 있으면 스킵
            if rag_sum_val and rag_det_val:
                sheet_skipped += 1
                continue

            question = str(q_val).strip()
            print(f"\n[{sheet_name} | 행 {r}] 질문: {question[:50]}...")

            t0 = time.time()
            try:
                # 1 & 2. RAG 검색 및 최종 답변 합성
                res, composed_answer = process_question_with_retry(ctx, question)

                # 3. 태그 파싱
                s_match = summary_pat.search(composed_answer)
                d_match = details_pat.search(composed_answer)

                if s_match and d_match:
                    summary = s_match.group(1).strip()
                    details = d_match.group(1).strip()
                elif s_match:
                    summary = s_match.group(1).strip()
                    details = ""
                else:
                    # 태그 매칭에 실패한 경우 전체를 요약에 기입
                    summary = composed_answer.strip()
                    details = ""
                    print("  -> [경고] XML 태그 파싱 실패! 전체 본문이 요약에 기입됩니다.")

                # 근거 파일 수집 (중복 제거 및 정렬)
                evidence = res.get("evidence") or []
                docs = sorted(list(set(ev.get("document_name") for ev in evidence if ev.get("document_name"))))
                evidence_files = ", ".join(docs)

                # 엑셀 셀에 값 쓰기
                sheet.cell(row=r, column=7, value=summary)
                sheet.cell(row=r, column=8, value=details)
                sheet.cell(row=r, column=9, value=evidence_files)

                # 파일 실시간 저장 (진척도 즉시 보존)
                wb.save(excel_path)

                elapsed = time.time() - t0
                print(f"  -> 완료! 소요시간: {elapsed:.2f}초")
                print(f"  -> 요약글자수: {len(summary)} / 상세글자수: {len(details)}")
                print(f"  -> 근거파일: {evidence_files if evidence_files else '(없음)'}")
                sheet_success += 1

                # API 제한 방지를 위한 대기
                time.sleep(3.5)

            except Exception as e:
                print(f"  -> [에러 발생] {e}")
                traceback.print_exc()
                wb.save(excel_path)
                # 에러 발생 시에도 중단하지 않고 잠시 대기 후 다음 행 진행할지 결정 가능하나, 심각한 문제 감지를 위해 로그 출력
                time.sleep(5.0)

        total_success += sheet_success
        total_skipped += sheet_skipped
        print(f"\n시트 [{sheet_name}] 완료: 새로 입력 {sheet_success}건, 스킵 {sheet_skipped}건")

    print("\n" + "="*50)
    print(f"전체 작업 완료! 총 새로 입력: {total_success}건, 총 스킵: {total_skipped}건")
    print(f"저장 경로: {excel_path}")
    print("="*50)
    return 0


if __name__ == "__main__":
    sys.exit(main())
